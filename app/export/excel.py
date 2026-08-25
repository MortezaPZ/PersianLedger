from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.domain import money


def _book() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.sheet_view.rightToLeft = True
    return wb


def _header_fill():
    return PatternFill("solid", fgColor="1E4D8C")


def _write_header(ws, row: int, values: list[str]) -> None:
    font = Font(name="Tahoma", bold=True, color="FFFFFF")
    fill = _header_fill()
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row, col, value)
        cell.font = font
        cell.fill = fill
        cell.alignment = align


def _autosize(ws) -> None:
    for column in ws.columns:
        letter = get_column_letter(column[0].column)
        width = 12
        for cell in column:
            if cell.value is None:
                continue
            width = max(width, min(len(str(cell.value)) + 2, 40))
        ws.column_dimensions[letter].width = width


def export_statement(path: Path, result) -> None:
    wb = _book()
    ws = wb.active
    ws.title = "صورتحساب"
    headers = ["تاریخ", "شماره", "نوع", "شرح", "سرفصل", "تفصیلی"]
    for col in result.currencies:
        headers.extend([f"{col.code} بدهکار", f"{col.code} بستانکار", f"{col.code} مانده"])
    _write_header(ws, 1, headers)
    for r, row in enumerate(result.rows, start=2):
        values = [row.date_j, row.number, row.doc_type, row.description, f"{row.account_code} {row.account_name}".strip(), row.tafsili]
        for col in result.currencies:
            debit, credit, balance = row.amounts.get(col.id, (0, 0, 0))
            values.extend([
                money.quantize(debit, col.decimals),
                money.quantize(credit, col.decimals),
                money.quantize(balance, col.decimals),
            ])
        for c, value in enumerate(values, start=1):
            ws.cell(r, c, value if value != "" else None)
    _autosize(ws)
    wb.save(path)


def export_trial_balance(path: Path, result) -> None:
    wb = _book()
    ws = wb.active
    ws.title = "تراز"
    labels = {
        2: ["مانده بد", "مانده بس"],
        4: ["گردش بد", "گردش بس", "مانده بد", "مانده بس"],
        6: ["افتتاحیه بد", "افتتاحیه بس", "گردش بد", "گردش بس", "مانده بد", "مانده بس"],
    }[result.columns]
    headers = ["کد", "عنوان"]
    for col in result.currencies:
        headers.extend([f"{col.code} {label}" for label in labels])
    _write_header(ws, 1, headers)
    for r, row in enumerate(result.rows, start=2):
        values = [row.code, row.name]
        for col in result.currencies:
            for offset in range(result.columns):
                values.append(row.cell(result.columns, offset, col.id))
        for c, value in enumerate(values, start=1):
            ws.cell(r, c, value)
    _autosize(ws)
    wb.save(path)


def export_journal(path: Path, result) -> None:
    wb = _book()
    ws = wb.active
    ws.title = "روزنامه"
    headers = ["تاریخ", "شماره", "ردیف", "سرفصل", "تفصیلی", "شرح"]
    for col in result.currencies:
        headers.extend([f"{col.code} بدهکار", f"{col.code} بستانکار"])
    _write_header(ws, 1, headers)
    for r, row in enumerate(result.rows, start=2):
        values = [row.date_j, row.number, row.line_no, row.account, row.tafsili, row.description]
        for col in result.currencies:
            debit, credit = row.amounts.get(col.id, (0, 0))
            values.extend([debit, credit])
        for c, value in enumerate(values, start=1):
            ws.cell(r, c, value if value != "" else None)
    _autosize(ws)
    wb.save(path)


def export_income_statement(path: Path, result) -> None:
    wb = _book()
    ws = wb.active
    ws.title = "سود و زیان"
    headers = ["کد", "عنوان"] + [col.code for col in result.currencies]
    _write_header(ws, 1, headers)
    row_no = 2

    def put(code, name, amounts, getter):
        nonlocal row_no
        values = [code, name] + [getter(amounts, col) for col in result.currencies]
        for c, value in enumerate(values, start=1):
            ws.cell(row_no, c, value)
        row_no += 1

    for line in result.income:
        put(line.code, line.name, line, lambda a, col: a.amount(col.id))
    put("", "جمع درآمدها", result.income_total, lambda a, col: a.get(col.id, Decimal(0)))
    for line in result.expense:
        put(line.code, line.name, line, lambda a, col: a.amount(col.id))
    put("", "جمع هزینه‌ها", result.expense_total, lambda a, col: a.get(col.id, Decimal(0)))
    put("", "سود (زیان) خالص", result.net, lambda a, col: a.get(col.id, Decimal(0)))
    _autosize(ws)
    wb.save(path)
